"""Extração do texto dos formatos do corpus: .docx, .xlsx e .pptx.

Cada formato vira uma sequência de `Unidade`s — o menor pedaço que o
chunking empacota — e os metadados que o próprio arquivo declara
(título, tabela de cabeçalho, classificação, perfis). Nada aqui decide o
que o arquivo É: isso é `identificacao.py`, que lê o que esta etapa
devolve.

Três regras valem para os três formatos:

- **o cabeçalho institucional e o rodapé saem do texto.** "Companhia do
  Metropolitano…" e "Documento gerado para ambiente de teste…" aparecem
  em todos os arquivos; indexados, fariam todo trecho parecer com todo
  trecho. A classificação e os perfis que vivem nessas linhas viram
  metadado, que é onde têm uso;
- **linha de tabela é indivisível.** Serializada como `coluna: valor`,
  ela carrega o próprio contexto; cortada ao meio, a segunda metade vira
  números sem nome;
- **célula vazia é omitida.** Os campos em branco plantados na base são
  assunto da base estruturada (onda C), que os encontra com `IS NULL`; no
  índice, `COLUNA: —` repetido em milhares de linhas só dilui a busca.
"""

import io
import re
from collections.abc import Iterator, Mapping
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any, Literal

from app.rag.ingestao.fonte import MIME_DOCX, MIME_PPTX, MIME_XLSX

Formato = Literal["docx", "xlsx", "pptx"]

_POR_MIME: dict[str, Formato] = {MIME_DOCX: "docx", MIME_XLSX: "xlsx", MIME_PPTX: "pptx"}

# "PMO Corporativo (GPM) · Classificação: Interno". No .pptx a mesma
# informação divide a linha com os perfis, separada por "·".
_CLASSIFICACAO = re.compile(r"Classifica[çc][ãa]o:\s*([^·|\n]+)", re.IGNORECASE)
_PERFIS = re.compile(r"Perfis autorizados:\s*(.+)$", re.IGNORECASE)
_INSTITUCIONAL = re.compile(r"^(Companhia do Metropolitano|PMO Corporativo)", re.IGNORECASE)
_RODAPE = re.compile(
    r"ambiente de teste|dados sint[ée]ticos de|integralmente sint[ée]tico", re.IGNORECASE
)
# Seção de primeiro nível: "5. ESCOPO PRELIMINAR". Subseção ("6.1 Registro
# da demanda") fica dentro da seção, como texto: separá-la produziria
# trechos de uma linha só.
_SECAO = re.compile(r"^\d{1,2}\.\s+\S")
_NUMERO_DE_SLIDE = re.compile(r"^\d{1,3}$")

GRUPO_CABECALHO = "Cabeçalho do documento"


class ErroExtracao(ValueError):
    """Arquivo num formato suportado, mas ilegível."""


class FormatoNaoSuportado(ErroExtracao):
    """Formato fora do escopo da ingestão."""


@dataclass(frozen=True)
class Unidade:
    """O menor pedaço de texto que o chunking empacota."""

    texto: str
    # Seção (.docx), aba (.xlsx) ou slide (.pptx). Vai para a linha de
    # contexto de cada trecho, e um trecho nunca atravessa dois grupos.
    grupo: str = ""
    # Linha de tabela ou de planilha: o chunking não a corta ao meio.
    indivisivel: bool = False
    # Só em linha de planilha: coluna → valor, já formatado. É daqui que a
    # identificação lê `ID_PROJETO` e as permissões leem a classificação
    # de cada linha do consolidado do portfólio.
    campos: Mapping[str, str] | None = None


@dataclass(frozen=True)
class DocumentoExtraido:
    formato: Formato
    # Linha de título (.docx) ou primeira linha do slide de capa (.pptx).
    titulo: str | None = None
    # Tabela de cabeçalho do .docx: Projeto, Código, Versão, Data...
    cabecalho: Mapping[str, str] = field(default_factory=dict)
    classificacao: str | None = None
    perfis_declarados: tuple[str, ...] = ()
    abas: tuple[str, ...] = ()
    unidades: tuple[Unidade, ...] = ()


def formato_de(nome: str, mime: str = "") -> Formato:
    """Formato pelo tipo MIME da fonte, com a extensão como reserva."""
    if mime in _POR_MIME:
        return _POR_MIME[mime]
    extensao = nome.rsplit(".", 1)[-1].lower() if "." in nome else ""
    if extensao in ("docx", "xlsx", "pptx"):
        return extensao  # type: ignore[return-value]
    raise FormatoNaoSuportado(
        f"formato fora do escopo da ingestão ({mime or extensao or 'sem extensão'})"
    )


def extrair(dados: bytes, nome: str, mime: str = "") -> DocumentoExtraido:
    """Lê o arquivo e devolve unidades e metadados declarados."""
    formato = formato_de(nome, mime)
    leitor = {"docx": _docx, "xlsx": _xlsx, "pptx": _pptx}[formato]
    try:
        return leitor(dados)
    except ErroExtracao:
        raise
    # Cada biblioteca de formato tem as próprias exceções; para a ingestão
    # todas significam a mesma coisa: arquivo ilegível, com o motivo.
    except Exception as erro:
        raise ErroExtracao(f"arquivo .{formato} ilegível: {erro}") from erro


def formatar_valor(valor: Any) -> str:
    """Valor de célula em texto: data em dd/mm/aaaa, número sem ruído de float.

    `802142.4300000001` vira `802142.43`: o ruído de ponto flutuante que
    a planilha carrega não é informação, e no texto indexado seria um
    número que ninguém digitaria numa pergunta.
    """
    if valor is None:
        return ""
    if isinstance(valor, datetime):
        if (valor.hour, valor.minute, valor.second) == (0, 0, 0):
            return valor.strftime("%d/%m/%Y")
        return valor.strftime("%d/%m/%Y %H:%M")
    if isinstance(valor, date):
        return valor.strftime("%d/%m/%Y")
    if isinstance(valor, bool):
        return "sim" if valor else "não"
    if isinstance(valor, float):
        if valor.is_integer():
            return str(int(valor))
        return f"{valor:.2f}".rstrip("0").rstrip(".")
    return " ".join(str(valor).split())


def _boilerplate(texto: str, achados: dict[str, Any]) -> bool:
    """Captura classificação e perfis; diz se a linha deve sair do texto."""
    classificacao = _CLASSIFICACAO.search(texto)
    if classificacao and "classificacao" not in achados:
        achados["classificacao"] = classificacao.group(1).strip()
    perfis = _PERFIS.search(texto)
    if perfis and "perfis" not in achados:
        achados["perfis"] = tuple(p.strip() for p in perfis.group(1).split(",") if p.strip())
    return bool(classificacao or perfis or _INSTITUCIONAL.match(texto) or _RODAPE.search(texto))


def _unidades_de_tabela(linhas: list[list[str]], grupo: str) -> list[Unidade]:
    """Serializa uma tabela em linhas `coluna: valor`, uma unidade por linha.

    Tabela de duas colunas é lida como chave e valor ("Término: 30/06/2027",
    "Gerente de Projeto: Manter atualizados os artefatos"); as demais usam
    a primeira linha como cabeçalho.
    """
    if not linhas:
        return []
    if all(len(linha) == 2 for linha in linhas):
        return [
            Unidade(f"{chave}: {valor}" if valor else chave, grupo, indivisivel=True)
            for chave, valor in linhas
        ]
    cabecalho, *corpo = linhas
    if not corpo:
        return [Unidade(" | ".join(c for c in cabecalho if c), grupo, indivisivel=True)]
    unidades = []
    for linha in corpo:
        pares = [f"{c}: {v}" if c else v for c, v in zip(cabecalho, linha) if v]
        if pares:
            unidades.append(Unidade(" | ".join(pares), grupo, indivisivel=True))
    return unidades


# --- .docx -------------------------------------------------------------------


def _docx(dados: bytes) -> DocumentoExtraido:
    from docx import Document
    from docx.table import Table

    documento = Document(io.BytesIO(dados))
    achados: dict[str, Any] = {}
    titulo: str | None = None
    cabecalho: dict[str, str] = {}
    unidades: list[Unidade] = []
    grupo = ""
    # Até a primeira seção, uma tabela de chave e valor é o cabeçalho do
    # documento (Projeto, Código, Versão, Data), e não conteúdo de seção.
    no_corpo = False

    for bloco in documento.iter_inner_content():
        if isinstance(bloco, Table):
            linhas = _linhas_docx(bloco)
            if not no_corpo and not cabecalho and _chave_valor(linhas):
                cabecalho = {chave: valor for chave, valor in linhas}
                unidades.extend(_unidades_de_tabela(linhas, GRUPO_CABECALHO))
                continue
            no_corpo = True
            unidades.extend(_unidades_de_tabela(linhas, grupo))
            continue

        texto = " ".join(bloco.text.split())
        if not texto or _boilerplate(texto, achados):
            continue
        if titulo is None and not no_corpo:
            titulo = texto
            continue
        no_corpo = True
        if _SECAO.match(texto) and texto == texto.upper():
            grupo = texto
            continue
        estilo = bloco.style.name if bloco.style is not None else ""
        if (estilo or "").lower().startswith("list"):
            texto = f"- {texto}"
        unidades.append(Unidade(texto, grupo))

    return DocumentoExtraido(
        formato="docx",
        titulo=titulo,
        cabecalho=cabecalho,
        classificacao=achados.get("classificacao"),
        perfis_declarados=achados.get("perfis", ()),
        unidades=tuple(unidades),
    )


def _linhas_docx(tabela: Any) -> list[list[str]]:
    """Linhas da tabela como listas de texto, sem repetir célula mesclada.

    O python-docx devolve a mesma célula uma vez por coluna da grade que
    ela ocupa; sem a deduplicação, uma célula mesclada em três colunas
    apareceria três vezes na linha serializada.
    """
    linhas = []
    for linha in tabela.rows:
        vistos: list[Any] = []
        celulas = []
        for celula in linha.cells:
            if any(celula._tc is visto for visto in vistos):
                continue
            vistos.append(celula._tc)
            celulas.append(" ".join(celula.text.split()))
        if any(celulas):
            linhas.append(celulas)
    return linhas


def _chave_valor(linhas: list[list[str]]) -> bool:
    return bool(linhas) and all(len(linha) == 2 and linha[0] for linha in linhas)


# --- .xlsx -------------------------------------------------------------------


def _xlsx(dados: bytes) -> DocumentoExtraido:
    from openpyxl import load_workbook

    livro = load_workbook(io.BytesIO(dados), read_only=True, data_only=True)
    abas: list[str] = []
    unidades: list[Unidade] = []
    try:
        for planilha in livro.worksheets:
            abas.append(planilha.title)
            colunas: list[str] | None = None
            for linha in planilha.iter_rows(values_only=True):
                valores = [formatar_valor(valor) for valor in linha]
                if colunas is None:
                    if any(valores):
                        colunas = [v or f"COLUNA_{i}" for i, v in enumerate(valores, start=1)]
                    continue
                campos = {coluna: valor for coluna, valor in zip(colunas, valores) if valor}
                if campos:
                    unidades.append(
                        Unidade(
                            " | ".join(f"{coluna}: {valor}" for coluna, valor in campos.items()),
                            f"aba {planilha.title}",
                            indivisivel=True,
                            campos=campos,
                        )
                    )
    finally:
        livro.close()
    return DocumentoExtraido(formato="xlsx", abas=tuple(abas), unidades=tuple(unidades))


# --- .pptx -------------------------------------------------------------------


def _pptx(dados: bytes) -> DocumentoExtraido:
    from pptx import Presentation

    apresentacao = Presentation(io.BytesIO(dados))
    achados: dict[str, Any] = {}
    titulo: str | None = None
    unidades: list[Unidade] = []

    for numero, slide in enumerate(apresentacao.slides, start=1):
        titulo_slide: str | None = None
        do_slide: list[Unidade] = []
        for forma in _formas(slide.shapes):
            if forma.has_chart:
                do_slide.extend(_unidades_de_grafico(forma.chart))
            elif forma.has_table:
                linhas = [
                    [" ".join(celula.text.split()) for celula in linha.cells]
                    for linha in forma.table.rows
                ]
                do_slide.extend(_unidades_de_tabela([li for li in linhas if any(li)], ""))
            elif forma.has_text_frame:
                for paragrafo in forma.text_frame.paragraphs:
                    texto = " ".join(paragrafo.text.split())
                    if not texto or _NUMERO_DE_SLIDE.match(texto) or _boilerplate(texto, achados):
                        continue
                    if titulo is None:
                        # Primeira linha da capa: o título do documento.
                        titulo = texto
                    elif titulo_slide is None and numero > 1:
                        titulo_slide = texto
                    else:
                        do_slide.append(Unidade(texto))
        grupo = f"slide {numero}" + (f": {titulo_slide}" if titulo_slide else "")
        unidades.extend(Unidade(u.texto, grupo, u.indivisivel, u.campos) for u in do_slide)

    return DocumentoExtraido(
        formato="pptx",
        titulo=titulo,
        classificacao=achados.get("classificacao"),
        perfis_declarados=achados.get("perfis", ()),
        unidades=tuple(unidades),
    )


def _formas(formas: Any) -> Iterator[Any]:
    """Formas do slide em ordem, abrindo os grupos."""
    from pptx.shapes.group import GroupShape

    for forma in formas:
        if isinstance(forma, GroupShape):
            yield from _formas(forma.shapes)
        else:
            yield forma


def _unidades_de_grafico(grafico: Any) -> list[Unidade]:
    """Uma unidade por série, com os valores que o gráfico guarda em cache.

    Os slides de visão geral do balanço são SÓ gráfico: sem esta leitura,
    a distribuição por fase e por farol não existiria no índice — e é nela
    que algumas das inconsistências plantadas aparecem, com o gráfico
    contradizendo o sumário executivo.
    """
    titulo = ""
    if grafico.has_title and grafico.chart_title.has_text_frame:
        titulo = " ".join(grafico.chart_title.text_frame.text.split())
    unidades = []
    for plot in grafico.plots:
        categorias = [str(categoria) for categoria in plot.categories]
        for serie in plot.series:
            pares = [
                f"{categoria}: {formatar_valor(valor)}"
                for categoria, valor in zip(categorias, serie.values)
                if valor is not None
            ]
            if pares:
                rotulo = " — ".join(parte for parte in ("Gráfico", titulo, serie.name) if parte)
                unidades.append(Unidade(f"{rotulo}: " + "; ".join(pares), indivisivel=True))
    return unidades
