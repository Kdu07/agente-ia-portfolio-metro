"""Testes da extração dos formatos do corpus (rag.ingestao.extracao).

Os arquivos são gerados em memória com as mesmas bibliotecas que os
leem, reproduzindo a forma dos documentos do parceiro: cabeçalho
institucional, linha de perfis, tabela de cabeçalho, seções numeradas,
rodapé de ambiente de teste. Nenhum binário versionado — e os testes
pulam limpos se o extra `rag` não estiver instalado.
"""

import io
from datetime import date, datetime

import pytest

from app.rag.ingestao.extracao import (
    GRUPO_CABECALHO,
    ErroExtracao,
    FormatoNaoSuportado,
    extrair,
    formatar_valor,
)


def _docx() -> bytes:
    docx = pytest.importorskip("docx")
    documento = docx.Document()
    for linha in (
        "COMPANHIA DO METROPOLITANO DE SÃO PAULO — METRÔ",
        "PMO Corporativo (GPM)  ·  Classificação: Interno",
        "Perfis autorizados: Diretoria, GPM, Gerente de Projeto, Equipe",
        "TERMO DE ABERTURA DE PROJETO",
    ):
        documento.add_paragraph(linha)
    cabecalho = documento.add_table(rows=2, cols=2)
    for posicao, (chave, valor) in enumerate((("Código", "PRJ-2024-018"), ("Versão", "v3"))):
        cabecalho.cell(posicao, 0).text = chave
        cabecalho.cell(posicao, 1).text = valor
    documento.add_paragraph("5. ESCOPO PRELIMINAR")
    documento.add_paragraph("Substituição integral do sistema de sinalização", style="List Bullet")
    documento.add_paragraph("6. ENTREGAS E MARCOS PRINCIPAIS")
    marcos = documento.add_table(rows=2, cols=3)
    colunas = (
        ("ID do marco", "MRC-PRJ-2024-018-01"),
        ("Descrição", "Aprovação da arquitetura"),
        ("Data prevista", "30/06/2024"),
    )
    for coluna, (titulo, valor) in enumerate(colunas):
        marcos.cell(0, coluna).text = titulo
        marcos.cell(1, coluna).text = valor
    documento.add_paragraph("6.1 Detalhamento das entregas")
    documento.add_paragraph(
        "Documento gerado para ambiente de teste. Conteúdo integralmente sintético."
    )
    destino = io.BytesIO()
    documento.save(destino)
    return destino.getvalue()


def _xlsx() -> bytes:
    openpyxl = pytest.importorskip("openpyxl")
    livro = openpyxl.Workbook()
    planilha = livro.active
    planilha.title = "Riscos e Problemas"
    planilha.append(["ID_ITEM", "ID_PROJETO", "PRAZO_ACAO", "IMPACTO_CUSTO_RS", "OBSERVACAO"])
    # Data sem fuso de propósito: é assim que a planilha a guarda.
    prazo = datetime(2026, 3, 23)  # noqa: DTZ001
    planilha.append(["RSK-PRJ-2024-018-001", "PRJ-2024-018", prazo, 802142.4300000001, None])
    destino = io.BytesIO()
    livro.save(destino)
    return destino.getvalue()


def _pptx() -> bytes:
    pytest.importorskip("pptx")
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches

    apresentacao = Presentation()
    branco = apresentacao.slide_layouts[6]

    capa = apresentacao.slides.add_slide(branco)
    quadro = capa.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2)).text_frame
    quadro.text = "Balanço do Portfólio — junho/2026"
    for linha in (
        "Companhia do Metropolitano de São Paulo — Metrô",
        "Classificação: Corporativo · Perfis autorizados: Diretoria, GPM, Gerente de Projeto",
        "1",
    ):
        quadro.add_paragraph().text = linha

    slide = apresentacao.slides.add_slide(branco)
    quadro = slide.shapes.add_textbox(Inches(1), Inches(0.5), Inches(8), Inches(1)).text_frame
    quadro.text = "Visão geral do portfólio"
    quadro.add_paragraph().text = "Distribuição por fase"
    dados = CategoryChartData()
    dados.categories = ["Demanda", "Execução"]
    dados.add_series("Projetos", (1, 9))
    slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1.5), Inches(6), Inches(3), dados
    )
    tabela = slide.shapes.add_table(2, 3, Inches(1), Inches(5), Inches(6), Inches(1)).table
    for coluna, (titulo, valor) in enumerate(
        (("ID", "PRJ-2024-018"), ("Projeto", "Modernização do Trecho Sul"), ("Farol", "Amarelo"))
    ):
        tabela.cell(0, coluna).text = titulo
        tabela.cell(1, coluna).text = valor
    rodape = slide.shapes.add_textbox(Inches(1), Inches(6.5), Inches(8), Inches(0.5)).text_frame
    rodape.text = "Balanço do Portfólio — junho/2026 · GPM · Classificação: Corporativo · Dados sintéticos de teste"
    slide.shapes.add_textbox(Inches(9), Inches(6.5), Inches(0.5), Inches(0.5)).text_frame.text = "2"

    destino = io.BytesIO()
    apresentacao.save(destino)
    return destino.getvalue()


# --- .docx ---------------------------------------------------------------------


def test_docx_separa_metadado_declarado_de_conteudo() -> None:
    documento = extrair(_docx(), "PRJ-2024-018_TAP_v3_20250822.docx")

    assert documento.titulo == "TERMO DE ABERTURA DE PROJETO"
    assert documento.cabecalho == {"Código": "PRJ-2024-018", "Versão": "v3"}
    assert documento.classificacao == "Interno"
    assert documento.perfis_declarados == ("Diretoria", "GPM", "Gerente de Projeto", "Equipe")


def test_docx_agrupa_por_secao_e_serializa_tabelas() -> None:
    grupos = {u.texto: u.grupo for u in extrair(_docx(), "x.docx").unidades}

    assert grupos["Código: PRJ-2024-018"] == GRUPO_CABECALHO
    assert grupos["- Substituição integral do sistema de sinalização"] == "5. ESCOPO PRELIMINAR"
    linha = (
        "ID do marco: MRC-PRJ-2024-018-01 | Descrição: Aprovação da arquitetura"
        " | Data prevista: 30/06/2024"
    )
    assert grupos[linha] == "6. ENTREGAS E MARCOS PRINCIPAIS"
    # Subseção não abre grupo: fica como texto da seção de primeiro nível.
    assert grupos["6.1 Detalhamento das entregas"] == "6. ENTREGAS E MARCOS PRINCIPAIS"


def test_docx_tira_o_boilerplate_do_texto() -> None:
    textos = [u.texto for u in extrair(_docx(), "x.docx").unidades]

    assert not any("COMPANHIA" in t or "ambiente de teste" in t or "Perfis" in t for t in textos)


# --- .xlsx ---------------------------------------------------------------------


def test_xlsx_serializa_cada_linha_com_os_nomes_das_colunas() -> None:
    documento = extrair(_xlsx(), "PRJ-2024-018_RSK_v1_20260811.xlsx")

    assert documento.abas == ("Riscos e Problemas",)
    (linha,) = documento.unidades
    assert linha.texto == (
        "ID_ITEM: RSK-PRJ-2024-018-001 | ID_PROJETO: PRJ-2024-018 | PRAZO_ACAO: 23/03/2026"
        " | IMPACTO_CUSTO_RS: 802142.43"
    )
    assert linha.indivisivel
    assert linha.grupo == "aba Riscos e Problemas"
    assert linha.campos["ID_PROJETO"] == "PRJ-2024-018"
    assert "OBSERVACAO" not in linha.campos  # célula vazia é omitida


# --- .pptx ---------------------------------------------------------------------


def test_pptx_le_capa_graficos_e_tabelas() -> None:
    documento = extrair(_pptx(), "PORT-2026_BAL_v1_20260630.pptx")
    textos = [u.texto for u in documento.unidades]

    assert documento.titulo == "Balanço do Portfólio — junho/2026"
    assert documento.classificacao == "Corporativo"
    assert documento.perfis_declarados == ("Diretoria", "GPM", "Gerente de Projeto")
    assert "Distribuição por fase" in textos
    assert any("Projetos: Demanda: 1; Execução: 9" in t for t in textos)
    assert "ID: PRJ-2024-018 | Projeto: Modernização do Trecho Sul | Farol: Amarelo" in textos
    assert {u.grupo for u in documento.unidades} == {"slide 2: Visão geral do portfólio"}
    assert not any("sintéticos" in t or t == "2" for t in textos)


# --- formatos e valores ------------------------------------------------------------


@pytest.mark.parametrize(
    ("valor", "esperado"),
    [
        (None, ""),
        (datetime(2026, 3, 23), "23/03/2026"),  # noqa: DTZ001 — célula de planilha não tem fuso
        (datetime(2026, 3, 23, 14, 5), "23/03/2026 14:05"),  # noqa: DTZ001
        (date(2025, 8, 22), "22/08/2025"),
        (802142.4300000001, "802142.43"),
        (94.40000000000001, "94.4"),
        (15.0, "15"),
        (True, "sim"),
        (4907, "4907"),
        ("  Em   tratamento ", "Em tratamento"),
    ],
)
def test_formatar_valor(valor, esperado) -> None:
    assert formatar_valor(valor) == esperado


def test_formato_fora_do_escopo_e_recusado() -> None:
    with pytest.raises(FormatoNaoSuportado):
        extrair(b"AVISO", "LEIA-ME_dados_sinteticos.txt", "text/plain")


def test_arquivo_corrompido_e_ilegivel_e_nao_fora_do_escopo() -> None:
    """A diferença vai para o registro: ignorado não é o mesmo que ilegível."""
    with pytest.raises(ErroExtracao) as erro:
        extrair(b"isto nao e um zip", "PRJ-2024-018_TAP_v3_20250822.docx")

    assert not isinstance(erro.value, FormatoNaoSuportado)
